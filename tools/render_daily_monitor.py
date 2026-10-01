#!/usr/bin/env python3
"""Render the daily monitoring workbook as a local, read-only HTML page.

The dashboard keeps the original XLSX unchanged and uses this HTML only for
in-page viewing.  It mirrors cell values, merged cells, dimensions and visible
formatting so the public page does not depend on an external office viewer.
"""

from __future__ import annotations

import datetime as dt
import html
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def theme_palette(workbook: openpyxl.Workbook) -> list[str]:
    """Return the Office theme colours in their indexed order."""
    default = [
        "FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD", "C0504D",
        "9BBB59", "8064A2", "4BACC6", "F79646", "0000FF", "800080",
    ]
    if not workbook.loaded_theme:
        return default

    try:
        root = ET.fromstring(workbook.loaded_theme)
        namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", namespace)
        if scheme is None:
            return default
        colours: list[str] = []
        for item in list(scheme):
            child = next(iter(item), None)
            value = child.attrib.get("lastClr") or child.attrib.get("val") if child is not None else None
            if not value:
                return default
            colours.append(value[-6:].upper())
        # OOXML stores the scheme as dk1, lt1, dk2, lt2; spreadsheet theme
        # indexes use lt1, dk1, lt2, dk2, then the accents.
        if len(colours) < 10:
            return default
        return [colours[1], colours[0], colours[3], colours[2], *colours[4:]]
    except (ET.ParseError, TypeError, AttributeError):
        return default


def tinted(hex_colour: str, tint: float | None) -> str:
    if tint in (None, 0):
        return hex_colour
    channels = []
    for index in (0, 2, 4):
        value = int(hex_colour[index:index + 2], 16)
        value = value * (1 + tint) if tint < 0 else value * (1 - tint) + 255 * tint
        channels.append(max(0, min(255, round(value))))
    return "".join(f"{value:02X}" for value in channels)


def colour(value, palette: list[str]) -> str | None:
    if value is None or value.type in (None, "auto"):
        return None
    if value.type == "rgb" and value.rgb:
        return value.rgb[-6:].upper()
    if value.type == "theme" and value.theme is not None and value.theme < len(palette):
        return tinted(palette[value.theme], value.tint)
    if value.type == "indexed" and value.indexed is not None:
        try:
            return openpyxl.styles.colors.COLOR_INDEX[value.indexed][-6:].upper()
        except IndexError:
            return None
    return None


def border_css(side, palette: list[str]) -> str:
    if side is None or side.style is None:
        return "none"
    width = {"thin": "1px", "medium": "2px", "thick": "3px"}.get(side.style, "1px")
    style = {"dashed": "dashed", "dotted": "dotted", "double": "double"}.get(side.style, "solid")
    return f"{width} {style} #{colour(side.color, palette) or '000000'}"


def value_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.strftime("%d.%m.%y")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return "\n".join(line.rstrip() for line in str(value).splitlines())


def cell_css(cell, palette: list[str]) -> str:
    font_colour = colour(cell.font.color, palette)
    fill_colour = colour(cell.fill.fgColor, palette) if cell.fill.fill_type else None
    align = cell.alignment
    font_size = cell.font.sz or 11
    css = [
        f"font-family:{html.escape(cell.font.name or 'Arial')},Arial,sans-serif",
        f"font-size:{font_size:g}pt",
        f"font-weight:{'700' if cell.font.bold else '400'}",
        f"font-style:{'italic' if cell.font.italic else 'normal'}",
        f"text-decoration:{'underline' if cell.font.underline else 'none'}",
        f"text-align:{align.horizontal or 'left'}",
        f"vertical-align:{align.vertical or 'middle'}",
        f"white-space:{'pre-wrap' if align.wrap_text else 'normal'}",
        f"border-top:{border_css(cell.border.top, palette)}",
        f"border-right:{border_css(cell.border.right, palette)}",
        f"border-bottom:{border_css(cell.border.bottom, palette)}",
        f"border-left:{border_css(cell.border.left, palette)}",
    ]
    if font_colour:
        css.append(f"color:#{font_colour}")
    if fill_colour:
        css.append(f"background-color:#{fill_colour}")
    if align.text_rotation:
        css.append(f"writing-mode:{'vertical-rl' if align.text_rotation in (90, 180) else 'horizontal-tb'}")
    return ";".join(css)


def render(source: Path, destination: Path) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True)
    worksheet = workbook.active
    palette = theme_palette(workbook)
    merged_starts: dict[tuple[int, int], tuple[int, int]] = {}
    merged_children: set[tuple[int, int]] = set()
    for area in worksheet.merged_cells.ranges:
        merged_starts[(area.min_row, area.min_col)] = (area.max_row - area.min_row + 1, area.max_col - area.min_col + 1)
        for row in range(area.min_row, area.max_row + 1):
            for column in range(area.min_col, area.max_col + 1):
                if (row, column) != (area.min_row, area.min_col):
                    merged_children.add((row, column))

    colgroup = []
    for column in range(1, worksheet.max_column + 1):
        width = worksheet.column_dimensions[get_column_letter(column)].width or 8.43
        colgroup.append(f'<col style="width:{max(34, round(width * 7.1))}px">')

    table_rows = []
    for row in range(1, worksheet.max_row + 1):
        height = worksheet.row_dimensions[row].height or 15
        cells = []
        for column in range(1, worksheet.max_column + 1):
            if (row, column) in merged_children:
                continue
            cell = worksheet.cell(row, column)
            rowspan, colspan = merged_starts.get((row, column), (1, 1))
            attrs = []
            if rowspan > 1:
                attrs.append(f'rowspan="{rowspan}"')
            if colspan > 1:
                attrs.append(f'colspan="{colspan}"')
            cells.append(
                f'<td {" ".join(attrs)} style="{cell_css(cell, palette)}">{html.escape(value_text(cell))}</td>'
            )
        table_rows.append(f'<tr style="height:{height}pt">{"".join(cells)}</tr>')

    document = f"""<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Ежедневный мониторинг</title>
  <style>
    html, body {{ margin:0; min-width:max-content; background:#fff; }}
    body {{ padding:12px; }}
    table {{ border-collapse:collapse; table-layout:fixed; background:#fff; }}
    td {{ box-sizing:border-box; min-width:34px; padding:2px 4px; line-height:1.2; overflow-wrap:anywhere; }}
  </style>
</head>
<body><table><colgroup>{''.join(colgroup)}</colgroup><tbody>{''.join(table_rows)}</tbody></table></body>
</html>"""
    destination.write_text(document, encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: render_daily_monitor.py INPUT.xlsx OUTPUT.html")
    render(Path(sys.argv[1]), Path(sys.argv[2]))
